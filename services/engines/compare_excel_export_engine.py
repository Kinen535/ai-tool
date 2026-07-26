from __future__ import annotations

import json
import re
from datetime import datetime
from io import BytesIO
from typing import Any, Mapping, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


TITLE_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FILL = PatternFill("solid", fgColor="D9EAF7")
GOOD_FILL = PatternFill("solid", fgColor="E2F0D9")
BAD_FILL = PatternFill("solid", fgColor="FCE4D6")
EXEMPT_FILL = PatternFill("solid", fgColor="DDEBF7")
WARNING_FILL = PatternFill("solid", fgColor="FFF2CC")
MUTED_FILL = PatternFill("solid", fgColor="E7E6E6")

THIN_BORDER = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
    bottom=Side(style="thin", color="D9D9D9"),
)

STATUS_FIELDS = (
    "综合考勤",
    "考勤结论",
    "考勤状态",
    "考勤结果",
    "综合结果",
    "结果",
    "状态",
)

STATUS_FILLS = {
    "合格": GOOD_FILL,
    "优秀": GOOD_FILL,
    "核心": GOOD_FILL,
    "未达标": BAD_FILL,
    "不合格": BAD_FILL,
    "缺勤": BAD_FILL,
    "清理": BAD_FILL,
    "违规": BAD_FILL,
    "豁免": EXEMPT_FILL,
    "待改进": WARNING_FILL,
    "待警告": WARNING_FILL,
    "无基线": MUTED_FILL,
}


def _safe_text(value: Any) -> str:
    if value is None:
        return ""

    if isinstance(value, (dict, list, tuple, set)):
        return json.dumps(
            value,
            ensure_ascii=False,
            separators=(",", ":"),
        )

    return str(value)


def _excel_value(value: Any) -> Any:
    if value is None:
        return ""

    if isinstance(value, (str, int, float, bool)):
        return value

    return _safe_text(value)


def _scope_label(payload: Mapping[str, Any]) -> str:
    scope = payload.get("analysis_scope") or {}

    if isinstance(scope, Mapping):
        scope_type = str(
            scope.get("scope_type")
            or payload.get("analysis_scope_type")
            or "alliance"
        )
        scope_value = str(
            scope.get("scope_value")
            or payload.get("analysis_scope_value")
            or ""
        )
    else:
        scope_type = str(
            payload.get("analysis_scope_type")
            or "alliance"
        )
        scope_value = str(
            payload.get("analysis_scope_value")
            or ""
        )

    if scope_type == "group":
        return f"指定分组-{scope_value or '未命名'}"

    if scope_type == "member":
        return f"指定成员-{scope_value or '未命名'}"

    return "全同盟"


def _safe_filename_part(value: str) -> str:
    value = re.sub(
        r'[\\/:*?"<>|]+',
        "-",
        value,
    )
    value = re.sub(r"\s+", "-", value).strip("-")
    return value[:60] or "未命名"


def _status_value(row: Mapping[str, Any]) -> str:
    for field in STATUS_FIELDS:
        value = row.get(field)

        if value not in (None, ""):
            return str(value).strip()

    return ""


def _row_matches_status(
    row: Mapping[str, Any],
    filter_value: str,
) -> bool:
    normalized_filter = str(
        filter_value or ""
    ).strip()

    if (
        not normalized_filter
        or normalized_filter.casefold() == "all"
        or normalized_filter in {
            "全部",
            "全部状态",
        }
    ):
        return True

    target = normalized_filter
    status = _status_value(row)

    if status == target:
        return True

    return any(
        target == str(value).strip()
        for value in row.values()
        if value not in (None, "")
    )


def _sort_key(value: Any) -> tuple[int, Any]:
    if value is None or value == "":
        return (2, "")

    if isinstance(value, (int, float)):
        return (0, float(value))

    text = str(value).strip()

    try:
        return (0, float(text.replace(",", "")))
    except ValueError:
        return (1, text.casefold())


def apply_client_state(
    rows: Sequence[Mapping[str, Any]],
    state: Mapping[str, Any] | None = None,
) -> tuple[list[dict[str, Any]], list[str]]:
    state = dict(state or {})
    output = [dict(row) for row in rows]

    status_filter = str(
        state.get("status_filter") or ""
    ).strip()

    output = [
        row
        for row in output
        if _row_matches_status(row, status_filter)
    ]

    sort_column = str(
        state.get("sort_column") or ""
    ).strip()

    sort_direction = str(
        state.get("sort_direction") or "asc"
    ).strip().lower()

    if sort_column:
        output.sort(
            key=lambda row: _sort_key(
                row.get(sort_column)
            ),
            reverse=sort_direction == "desc",
        )

    requested_columns = state.get("visible_columns")
    columns: list[str] = []

    if isinstance(requested_columns, list):
        columns = [
            str(column)
            for column in requested_columns
            if str(column)
        ]

    if not columns:
        for row in output or rows:
            for column in row.keys():
                if column not in columns:
                    columns.append(str(column))

    return output, columns


def _flatten(
    value: Any,
    prefix: str = "",
) -> list[tuple[str, Any]]:
    rows: list[tuple[str, Any]] = []

    if isinstance(value, Mapping):
        for key, child in value.items():
            name = f"{prefix}.{key}" if prefix else str(key)
            rows.extend(_flatten(child, name))

        return rows

    if isinstance(value, (list, tuple)):
        if all(
            not isinstance(item, (dict, list, tuple))
            for item in value
        ):
            rows.append((prefix, "、".join(map(str, value))))
        else:
            rows.append((prefix, _safe_text(value)))

        return rows

    rows.append((prefix or "值", value))
    return rows


def _style_title(sheet, title: str, width: int = 2) -> None:
    sheet.insert_rows(1)
    sheet.merge_cells(
        start_row=1,
        start_column=1,
        end_row=1,
        end_column=max(2, width),
    )

    cell = sheet.cell(1, 1, title)
    cell.font = Font(
        bold=True,
        color="FFFFFF",
        size=14,
    )
    cell.fill = TITLE_FILL
    cell.alignment = Alignment(
        horizontal="center",
        vertical="center",
    )
    sheet.row_dimensions[1].height = 24


def _finish_sheet(sheet) -> None:
    sheet.sheet_view.showGridLines = False

    for column_cells in sheet.columns:
        values = [
            len(_safe_text(cell.value))
            for cell in column_cells
            if cell.value not in (None, "")
        ]

        width = max(values, default=8) + 2
        width = max(10, min(width, 42))

        sheet.column_dimensions[
            get_column_letter(column_cells[0].column)
        ].width = width


def _write_key_value_sheet(
    workbook: Workbook,
    name: str,
    title: str,
    value: Any,
) -> None:
    sheet = workbook.create_sheet(name)
    sheet.append(["项目", "内容"])

    for key, item in _flatten(value):
        sheet.append([key, _excel_value(item)])

    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.border = THIN_BORDER

    for row in sheet.iter_rows(
        min_row=2,
        max_row=sheet.max_row,
        min_col=1,
        max_col=2,
    ):
        for cell in row:
            cell.border = THIN_BORDER
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )

    sheet.freeze_panes = "A3"
    sheet.auto_filter.ref = (
        f"A2:B{max(sheet.max_row, 2)}"
    )

    _style_title(sheet, title, 2)
    _finish_sheet(sheet)


def _status_fill_for_row(
    row: Mapping[str, Any],
) -> PatternFill | None:
    status = _status_value(row)

    for keyword, fill in STATUS_FILLS.items():
        if keyword in status:
            return fill

    return None


def _write_rows_sheet(
    workbook: Workbook,
    name: str,
    title: str,
    rows: Sequence[Mapping[str, Any]],
    columns: Sequence[str] | None = None,
) -> None:
    sheet = workbook.create_sheet(name)
    row_list = [dict(row) for row in rows]

    resolved_columns = list(columns or [])

    if not resolved_columns:
        for row in row_list:
            for column in row.keys():
                if column not in resolved_columns:
                    resolved_columns.append(str(column))

    if not resolved_columns:
        resolved_columns = ["说明"]
        row_list = [{"说明": "暂无数据"}]

    sheet.append(resolved_columns)

    for row in row_list:
        sheet.append([
            _excel_value(row.get(column))
            for column in resolved_columns
        ])

    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.border = THIN_BORDER
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

    for index, row in enumerate(row_list, start=2):
        status_fill = _status_fill_for_row(row)

        for cell in sheet[index]:
            cell.border = THIN_BORDER
            cell.alignment = Alignment(
                vertical="center",
                wrap_text=True,
            )

            if status_fill is not None:
                cell.fill = status_fill

            header = str(
                sheet.cell(1, cell.column).value or ""
            )

            if (
                isinstance(cell.value, (int, float))
                and (
                    "率" in header
                    or "权重" in header
                )
                and 0 <= float(cell.value) <= 1
            ):
                cell.number_format = "0.0%"

            elif isinstance(cell.value, (int, float)):
                cell.number_format = "#,##0.00"

    sheet.freeze_panes = "A3"
    sheet.auto_filter.ref = (
        f"A2:{get_column_letter(sheet.max_column)}"
        f"{max(sheet.max_row, 2)}"
    )

    _style_title(
        sheet,
        title,
        max(2, sheet.max_column),
    )
    _finish_sheet(sheet)


def _advice_rows(
    advice: Mapping[str, Any],
) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []

    for category, members in advice.items():
        if not isinstance(members, list):
            continue

        if not members:
            rows.append({
                "名单类型": category,
                "成员": "暂无",
            })
            continue

        for member in members:
            rows.append({
                "名单类型": category,
                "成员": member,
            })

    return rows


def build_compare_excel(
    payload: Mapping[str, Any],
    *,
    mode: str = "smart_report",
    client_state: Mapping[str, Any] | None = None,
    exported_at: datetime | None = None,
) -> tuple[bytes, str]:
    exported_at = exported_at or datetime.now()
    mode = str(mode or "smart_report").strip()

    if mode not in {
        "smart_report",
        "current_view",
    }:
        raise ValueError(
            f"不支持的导出模式：{mode}"
        )

    all_rows = payload.get("all_data")

    if not isinstance(all_rows, list):
        all_rows = payload.get("data") or []

    all_rows = [
        dict(row)
        for row in all_rows
        if isinstance(row, Mapping)
    ]

    current_rows, current_columns = apply_client_state(
        all_rows,
        client_state,
    )

    attendance = payload.get("attendance") or {}
    groups = (
        attendance.get("分组概览")
        if isinstance(attendance, Mapping)
        else None
    )

    if not isinstance(groups, list):
        groups = payload.get("groups") or []

    advice = payload.get("advice") or {}

    scope_label = _scope_label(payload)

    metadata = {
        "导出模式": (
            "智能分析报告"
            if mode == "smart_report"
            else "当前筛选结果"
        ),
        "导出时间": exported_at.strftime(
            "%Y-%m-%d %H:%M:%S"
        ),
        "分析范围": scope_label,
        "开始快照": payload.get("selected_old", ""),
        "结束快照": payload.get("selected_new", ""),
        "分组关键词": payload.get("team_keyword", ""),
        "成员关键词": payload.get("member_keyword", ""),
        "势力增长下限": payload.get(
            "power_growth_min",
            "",
        ),
        "势力增长上限": payload.get(
            "power_growth_max",
            "",
        ),
        "战功增长下限": payload.get("war_min", ""),
        "战功增长上限": payload.get("war_max", ""),
        "助攻增长下限": payload.get(
            "assist_min",
            "",
        ),
        "助攻增长上限": payload.get(
            "assist_max",
            "",
        ),
        "捐献增长下限": payload.get(
            "donate_min",
            "",
        ),
        "捐献增长上限": payload.get(
            "donate_max",
            "",
        ),
        "后台结果人数": len(all_rows),
        "当前视图人数": len(current_rows),
        "前端快捷筛选": (
            dict(client_state or {})
        ),
    }

    workbook = Workbook()
    workbook.remove(workbook.active)

    _write_key_value_sheet(
        workbook,
        "导出说明",
        "对比分析Excel导出说明",
        metadata,
    )

    if mode == "current_view":
        _write_rows_sheet(
            workbook,
            "当前筛选结果",
            "当前筛选结果",
            current_rows,
            current_columns,
        )
    else:
        overview = {}

        if isinstance(attendance, Mapping):
            overview = {
                "分析范围": attendance.get("范围", {}),
                "统计周期": attendance.get("统计周期", {}),
                "同盟概览": attendance.get(
                    "同盟概览",
                    {},
                ),
                "配置状态": attendance.get(
                    "配置状态",
                    "",
                ),
                "配置来源": attendance.get(
                    "配置来源",
                    {},
                ),
                "豁免规则": attendance.get(
                    "豁免规则",
                    "",
                ),
                "豁免成员数": attendance.get(
                    "豁免成员数",
                    0,
                ),
            }

        _write_key_value_sheet(
            workbook,
            "驾驶舱总览",
            "对比分析驾驶舱总览",
            overview,
        )

        _write_rows_sheet(
            workbook,
            "分组考勤概览",
            "分组考勤概览",
            [
                dict(row)
                for row in groups
                if isinstance(row, Mapping)
            ],
        )

        _write_rows_sheet(
            workbook,
            "成员明细",
            "完整成员考勤明细",
            all_rows,
        )

        _write_rows_sheet(
            workbook,
            "执行名单",
            "管理执行名单",
            _advice_rows(
                advice
                if isinstance(advice, Mapping)
                else {}
            ),
            ["名单类型", "成员"],
        )

        config = {}

        if isinstance(attendance, Mapping):
            config = attendance.get("配置") or {}

        _write_key_value_sheet(
            workbook,
            "考勤标准",
            "考勤阈值与有效权重",
            {
                "考勤配置": config,
                "页面配置": payload.get(
                    "attendance_config",
                    {},
                ),
            },
        )

    output = BytesIO()
    workbook.save(output)

    scope_part = _safe_filename_part(scope_label)
    date_part = exported_at.strftime("%Y%m%d-%H%M%S")

    mode_part = (
        "智能分析报告"
        if mode == "smart_report"
        else "当前筛选结果"
    )

    filename = (
        f"同盟考勤_{scope_part}_"
        f"{mode_part}_{date_part}.xlsx"
    )

    return output.getvalue(), filename
