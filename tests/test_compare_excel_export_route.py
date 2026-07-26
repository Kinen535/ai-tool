from __future__ import annotations

import json
import sqlite3
from io import BytesIO

import pytest
from openpyxl import load_workbook
from werkzeug.exceptions import BadRequest

import app as app_module


@pytest.fixture(autouse=True)
def use_temporary_database(tmp_path, monkeypatch):
    monkeypatch.setattr(
        app_module,
        "DB_FILE",
        tmp_path / "snapshots.db",
    )


PAYLOAD = {
    "selected_old": "2026-07-26 08:00:00",
    "selected_new": "2026-07-27 08:00:00",
    "analysis_scope": {
        "scope_type": "group",
        "scope_value": "折柳一组",
    },
    "all_data": [
        {
            "成员": "成员甲",
            "分组": "折柳一组",
            "战功增长": 5000,
            "考勤状态": "合格",
        },
        {
            "成员": "成员乙",
            "分组": "折柳一组",
            "战功增长": 500,
            "考勤状态": "未达标",
        },
    ],
    "attendance": {
        "范围": {"范围人数": 2},
        "配置": {"thresholds": {"battle": 2000}},
        "同盟概览": {"总人数": 2},
        "分组概览": [
            {"分组": "折柳一组", "人数": 2}
        ],
    },
    "advice": {
        "清理名单": ["成员乙"],
        "核心成员": ["成员甲"],
    },
}


def save_cache():
    conn = sqlite3.connect(app_module.DB_FILE)

    try:
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS battles (
                id INTEGER PRIMARY KEY,
                is_current INTEGER NOT NULL DEFAULT 0
            )
            """
        )
        conn.execute(
            "INSERT INTO battles (id, is_current) VALUES (1, 1)"
        )

        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS compare_cache (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                battle_id INTEGER,
                snapshot_time TEXT,
                data_json TEXT
            )
            """
        )

        conn.execute(
            "DELETE FROM compare_cache"
        )

        conn.execute(
            """
            INSERT INTO compare_cache (
                battle_id,
                snapshot_time,
                data_json
            )
            VALUES (?, ?, ?)
            """,
            (
                1,
                "excel-route-test",
                json.dumps(
                    PAYLOAD,
                    ensure_ascii=False,
                ),
            ),
        )

        conn.commit()
    finally:
        conn.close()


def call_route(body):
    with app_module.app.test_request_context(
        "/compare/export.xlsx",
        method="POST",
        json=body,
    ):
        response = (
            app_module.export_compare_excel_xlsx()
        )
        response.direct_passthrough = False
        return response


def test_smart_report_route():
    save_cache()

    response = call_route({
        "mode": "smart_report",
    })

    assert response.status_code == 200
    assert response.mimetype.endswith(
        "spreadsheetml.sheet"
    )

    workbook = load_workbook(
        BytesIO(response.get_data())
    )

    assert workbook.sheetnames == [
        "导出说明",
        "驾驶舱总览",
        "分组考勤概览",
        "成员明细",
        "执行名单",
        "考勤标准",
    ]


def test_current_view_route():
    save_cache()

    response = call_route({
        "mode": "current_view",
        "client_state": {
            "status_filter": "未达标",
            "visible_columns": [
                "成员",
                "考勤状态",
            ],
        },
    })

    workbook = load_workbook(
        BytesIO(response.get_data())
    )

    sheet = workbook["当前筛选结果"]

    assert sheet["A3"].value == "成员乙"
    assert sheet["B3"].value == "未达标"
    assert sheet.max_row == 3


def test_invalid_mode_returns_400():
    save_cache()

    with pytest.raises(BadRequest):
        call_route({
            "mode": "unknown",
        })
