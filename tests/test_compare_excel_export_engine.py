from __future__ import annotations

from datetime import datetime
from io import BytesIO

from openpyxl import load_workbook

from services.engines.compare_excel_export_engine import (
    apply_client_state,
    build_compare_excel,
)


def sample_payload():
    return {
        "selected_old": "2026-07-26 08:00:00",
        "selected_new": "2026-07-27 08:00:00",
        "analysis_scope": {
            "scope_type": "group",
            "scope_value": "折柳一组",
        },
        "team_keyword": "",
        "member_keyword": "",
        "all_data": [
            {
                "成员": "成员甲",
                "分组": "折柳一组",
                "战功增长": 5200,
                "助攻增长": 160,
                "考勤状态": "合格",
            },
            {
                "成员": "成员乙",
                "分组": "折柳一组",
                "战功增长": 1200,
                "助攻增长": 20,
                "考勤状态": "未达标",
            },
            {
                "成员": "成员丙",
                "分组": "折柳一组",
                "战功增长": 0,
                "助攻增长": 0,
                "考勤状态": "豁免",
            },
        ],
        "groups": [
            {
                "分组": "折柳一组",
                "人数": 3,
                "合格人数": 1,
            }
        ],
        "attendance": {
            "范围": {
                "分析范围": "指定分组",
                "范围人数": 3,
            },
            "配置": {
                "thresholds": {
                    "battle": 2000,
                    "assist": 100,
                    "donate": 100,
                },
                "weights": {
                    "battle": 0.6,
                    "assist": 0.4,
                    "donate": 0,
                },
            },
            "同盟概览": {
                "总人数": 3,
                "合格人数": 1,
                "合格率": 1 / 3,
            },
            "分组概览": [
                {
                    "分组": "折柳一组",
                    "人数": 3,
                    "合格人数": 1,
                    "合格率": 1 / 3,
                }
            ],
            "统计周期": {
                "开始": "2026-07-26 08:00:00",
                "结束": "2026-07-27 08:00:00",
            },
            "配置状态": "已启用",
            "豁免成员数": 1,
        },
        "advice": {
            "清理名单": ["成员乙"],
            "警告名单": [],
            "核心成员": ["成员甲"],
            "未执行名单": [],
        },
    }


def test_client_state_filters_sorts_and_columns():
    rows, columns = apply_client_state(
        sample_payload()["all_data"],
        {
            "status_filter": "合格",
            "sort_column": "战功增长",
            "sort_direction": "desc",
            "visible_columns": [
                "成员",
                "战功增长",
                "考勤状态",
            ],
        },
    )

    assert len(rows) == 1
    assert rows[0]["成员"] == "成员甲"
    assert columns == [
        "成员",
        "战功增长",
        "考勤状态",
    ]


def test_current_view_workbook():
    content, filename = build_compare_excel(
        sample_payload(),
        mode="current_view",
        client_state={
            "status_filter": "未达标",
            "visible_columns": [
                "成员",
                "分组",
                "考勤状态",
            ],
        },
        exported_at=datetime(
            2026,
            7,
            27,
            1,
            0,
            0,
        ),
    )

    workbook = load_workbook(BytesIO(content))

    assert workbook.sheetnames == [
        "导出说明",
        "当前筛选结果",
    ]

    sheet = workbook["当前筛选结果"]

    assert sheet.freeze_panes == "A3"
    assert sheet["A3"].value == "成员乙"
    assert sheet["C3"].value == "未达标"
    assert sheet.max_row == 3
    assert filename.endswith(".xlsx")
    assert "折柳一组" in filename


def test_smart_report_workbook():
    content, filename = build_compare_excel(
        sample_payload(),
        mode="smart_report",
        exported_at=datetime(
            2026,
            7,
            27,
            1,
            0,
            0,
        ),
    )

    workbook = load_workbook(BytesIO(content))

    assert workbook.sheetnames == [
        "导出说明",
        "驾驶舱总览",
        "分组考勤概览",
        "成员明细",
        "执行名单",
        "考勤标准",
    ]

    assert workbook["成员明细"].max_row == 5
    assert workbook["分组考勤概览"]["A3"].value == "折柳一组"
    assert workbook["执行名单"].max_row >= 4
    assert "智能分析报告" in filename
    assert len(content) > 5000



def test_status_filter_all_keeps_all_rows():
    from services.engines.compare_excel_export_engine import (
        apply_client_state,
    )

    rows = [
        {
            "成员": "成员甲",
            "状态": "核心",
        },
        {
            "成员": "成员乙",
            "状态": "待警告",
        },
    ]

    filtered, columns = apply_client_state(
        rows,
        {
            "status_filter": "all",
        },
    )

    assert [
        row["成员"]
        for row in filtered
    ] == [
        "成员甲",
        "成员乙",
    ]

    assert "成员" in columns
